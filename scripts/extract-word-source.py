from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
from PIL import Image
from lxml import etree
import json, re, argparse
parser=argparse.ArgumentParser(description='Extract profile data and photographs from the curated Word compilation.')
parser.add_argument('source_docx', help='Path to the curated Word compilation .docx')
args=parser.parse_args()
src=Path(args.source_docx).expanduser().resolve()
root=Path(__file__).resolve().parents[1]
assets=root/'public'/'assets'/'profiles'
assets.mkdir(parents=True,exist_ok=True)
doc=Document(src)
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
field_map={'Full Name':'name','Nick Name':'nickname','Branch':'branch','Email':'email','Phone':'phone','Current Location':'location','Educational Qualification':'education','Professional':'career','Family':'family','Memory from College Days':'memory'}

def form_row(t):
    # Row index where a "Full Name" form starts in this table, or None.
    # (A form may start partway down a table, e.g. below the previous person's photos.)
    for i,row in enumerate(t.rows):
        try:
            if row.cells[0].text.strip().replace('\n',' ')=='Full Name': return i
        except Exception:
            pass
    return None

def is_profile_table(t):
    return form_row(t) is not None

# Optional index table: only if the first table is NOT a profile form (older compilations had one)
index=[]
if doc.tables and not is_profile_table(doc.tables[0]):  # older compilations only
    for row in doc.tables[0].rows[1:]:
        vals=[c.text.strip() for c in row.cells]
        if len(vals)>=3 and vals[1]: index.append({'number':vals[0].strip('. '),'name':vals[1],'branch':vals[2]})

def collect_photos(rows, fields, slug, pdir, photos, seen):
    # Captions are often Word textboxes attached to the photo table, not ordinary table-cell text.
    table_captions=[]
    textboxes=[]
    for row in rows:
        textboxes+=etree.fromstring(row._tr.xml.encode('utf-8')).xpath('.//w:txbxContent',namespaces=NS) if hasattr(row._tr,'xml') else []
    for textbox in textboxes:
        cap=' '.join((node.text or '').strip() for node in textbox.xpath('.//w:t',namespaces=NS) if (node.text or '').strip())
        cap=' '.join(cap.split())
        if cap and cap not in table_captions: table_captions.append(cap)
    start=len(photos)
    for ri,row in enumerate(rows):
        for ci,cell in enumerate(row.cells):
            blips=cell._tc.xpath('.//a:blip')
            celltexts=cell._tc.xpath('.//w:txbxContent//w:t | .//w:p[not(ancestor::w:txbxContent)]//w:t')
            raw=' '.join((x.text or '').strip() for x in celltexts if (x.text or '').strip())
            bits=[]
            for bit in re.split(r'(?<=[.!?])\s+|\n+',raw):
                bit=bit.strip()
                if bit and bit not in bits: bits.append(bit)
            caption=' '.join(bits)
            words=caption.split()
            if len(words)%2==0 and words[:len(words)//2]==words[len(words)//2:]:
                caption=' '.join(words[:len(words)//2])
            for blip in blips:
                rid=blip.get(qn('r:embed'))
                if not rid or rid in seen: continue
                seen.add(rid)
                part=doc.part.related_parts.get(rid)
                if not part: continue
                ext=part.content_type.split('/')[-1].replace('jpeg','jpg')
                if ext not in ('jpg','png','gif','webp','bmp'): ext='jpg'
                fn=f'photo-{len(photos)+1}.{ext}'
                (pdir/fn).write_bytes(part.blob)
                cap=caption
                if cap.lower() in {fields['name'].lower()+': now', fields['name'].lower()+': then', fields['name'].lower()+' now', fields['name'].lower()+' then'}: cap=''
                photos.append({'src':f'/assets/profiles/{slug}/{fn}','caption':cap,'alt':fields['name']+' photograph '+str(len(photos)+1),'order':len(photos)+1})
    new=photos[start:]
    if len(table_captions)==len(new):
        for j,ph in enumerate(new):
            cap=table_captions[j]
            cap=cap.replace('P awan Jain','Pawan Jain').replace('P awan J ain','Pawan Jain').replace('S cene s','Scenes').replace('A lso','Also')
            ph['caption']=cap

profiles=[]
tables=doc.tables
ti=0
while ti<len(tables):
    t=tables[ti]
    fr=form_row(t)
    if fr is None:
        ti+=1; continue
    fields={}
    for row in t.rows[fr:]:
        if len(row.cells)>=2:
            k=row.cells[0].text.strip().replace('\n',' ')
            v=row.cells[1].text.strip()
            if k in field_map: fields[field_map[k]]=v
    if not fields.get('name'):
        ti+=1; continue
    slug=re.sub(r'[^a-z0-9]+','-',fields['name'].lower()).strip('-')
    pdir=assets/slug; pdir.mkdir(parents=True,exist_ok=True)
    photos=[]; seen=set()
    # photo rows that sit above the form in the same table belong to the PREVIOUS person (handled below);
    # every table after this one and before the next form is a photo table for this person
    pti=ti+1
    while pti<len(tables):
        nfr=form_row(tables[pti])
        if nfr is None:
            collect_photos(tables[pti].rows,fields,slug,pdir,photos,seen)
            pti+=1
        else:
            if nfr>0:  # photos stacked above the next person's form
                collect_photos(tables[pti].rows[:nfr],fields,slug,pdir,photos,seen)
            break
    fields['nickname']='' if fields.get('nickname','').strip() in ('--','-','—') else fields.get('nickname','').strip()
    fields['participant_id']=slug
    fields['photos']=photos
    profiles.append(fields)
    ti=pti
profiles.sort(key=lambda p:p['name'].casefold())
if not index:
    # No separate index table in this compilation: the contributor index is the list of profile names and branches from the Word forms.
    index=[{'number':str(i+1),'name':p['name'],'branch':p.get('branch','')} for i,p in enumerate(profiles)]
index_sorted=sorted(index,key=lambda p:p['name'].casefold())
payload={'book':{'title':'Golden Jubilee ’76','subtitle':'Golden Jubilee Celebration • NIT Raipur','years':'1976–2026','institution':'National Institute of Technology Raipur','website':'boroai.in'},'index':index_sorted,'participants':profiles,'sourceAudit':{'sourceDocument':src.name,'indexEntries':len(index),'profilesExtracted':len(profiles),'photographsExtracted':sum(len(p['photos']) for p in profiles),'indexOnlyEntries':[x for x in index_sorted if not any(p['name'].casefold()==x['name'].casefold() for p in profiles)]}}
(root/'public'/'public_data.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
(root/'public_data.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(payload['sourceAudit'],ensure_ascii=False,indent=2))
for p in profiles: print(p['name'],len(p['photos']),[(x['caption'][:45]) for x in p['photos']])
